import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def update_sih_presentation():
    prs = Presentation("SIH_Submission_StormBreaker.pptx")
    
    # ---------------- SLIDE 1: Title Page ----------------
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if shape.name == 'TextBox 9' and shape.has_text_frame:
            text_frame = shape.text_frame
            text_frame.clear()
            p = text_frame.paragraphs[0]
            p.text = "Problem Statement ID - [Leave Blank]"
            p.font.size = Pt(18)
            
            p = text_frame.add_paragraph()
            p.text = "Problem Statement Title - Detection & Characterization of Subsurface Ice in Lunar South Polar Doubly-Shadowed Craters"
            p.font.size = Pt(18)
            
            p = text_frame.add_paragraph()
            p.text = "Theme - Space Technology"
            p.font.size = Pt(18)
            
            p = text_frame.add_paragraph()
            p.text = "PS Category - Software"
            p.font.size = Pt(18)
            
            p = text_frame.add_paragraph()
            p.text = "Team ID - [Leave Blank]"
            p.font.size = Pt(18)
            
            p = text_frame.add_paragraph()
            p.text = "Team Name (Registered on portal) - storm Breaker"
            p.font.size = Pt(18)
            
            # Ensure text is black and readable
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    
    # Helper to update team name in ovals
    def update_team_name(slide):
        for shape in slide.shapes:
            if "Oval" in shape.name and shape.has_text_frame:
                shape.text_frame.text = "storm Breaker"
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.size = Pt(14)
                        r.font.bold = True

    # ---------------- SLIDE 2: Proposed Solution ----------------
    slide2 = prs.slides[1]
    update_team_name(slide2)
    slide2.shapes[1].text_frame.text = "IDEA TITLE: Lumina - Lunar Digital Twin"
    for shape in slide2.shapes:
        if shape.name == 'TextBox 8' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Proposed Solution:"
            p.font.bold = True
            p.font.size = Pt(20)
            
            p = tf.add_paragraph()
            p.text = "• A real-time 3D browser-native Digital Twin mirroring Faustini Crater using Chandrayaan-2 DFSAR and CLASS data."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Orchestrates 64 autonomous micro-rovers using Boids Swarm Algorithm for optimal sub-surface ice mapping."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Unique USP: Explainable AI reasoning cards and physics-accurate 2.6-second Earth-Moon latency simulator."
            p.font.size = Pt(18)
            
    img_path = "slide_6_7_extracted/screenshot2_missioncontrol.png"
    if os.path.exists(img_path):
        slide2.shapes.add_picture(img_path, Inches(0.5), Inches(3.5), width=Inches(8))

    # ---------------- SLIDE 3: Technical Approach ----------------
    slide3 = prs.slides[2]
    update_team_name(slide3)
    for shape in slide3.shapes:
        if shape.name == 'TextBox 8' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Core Tech Stack:"
            p.font.bold = True
            p.font.size = Pt(20)
            p = tf.add_paragraph()
            p.text = "• Frontend & 3D: Next.js 16, React 19, Three.js, React Three Fiber (WebGL 60FPS)."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• AI & Data: TensorFlow.js (NavCam CV), Zustand (State), Craig Reynolds' Boids Algorithm."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• No-Server Architecture: Runs 100% client-side via Vercel for zero infrastructure cost."
            p.font.size = Pt(18)
            
    img_path2 = "slide_6_7_extracted/lumina_architecture_diagram_1782737865178.png"
    if os.path.exists(img_path2):
        slide3.shapes.add_picture(img_path2, Inches(1.5), Inches(3.2), height=Inches(3.8))

    # ---------------- SLIDE 4: Feasibility and Viability ----------------
    slide4 = prs.slides[3]
    update_team_name(slide4)
    for shape in slide4.shapes:
        if shape.name == 'TextBox 8' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Feasibility & Risk Mitigation:"
            p.font.bold = True
            p.font.size = Pt(20)
            p = tf.add_paragraph()
            p.text = "• High Feasibility: Built entirely on open-source tools and publicly available ISRO/NASA datasets."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Challenge: Rendering 64 complex 3D rovers without lag."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Solution: Utilizes Three.js InstancedMesh to draw all rovers in a single GPU draw call."
            p.font.size = Pt(18)
            
    img_path3 = "slide_6_7_extracted/screenshot3_navcam.png"
    if os.path.exists(img_path3):
        slide4.shapes.add_picture(img_path3, Inches(0.5), Inches(3.2), width=Inches(8))

    # ---------------- SLIDE 5: Impact and Benefits ----------------
    slide5 = prs.slides[4]
    update_team_name(slide5)
    for shape in slide5.shapes:
        if shape.name == 'TextBox 8' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Strategic Impact for ISRO:"
            p.font.bold = True
            p.font.size = Pt(20)
            p = tf.add_paragraph()
            p.text = "• Pre-deployment Testing: Allows mission planners to simulate complex multi-agent swarms safely."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Explainable AI: Builds trust between human operators and AI by visualizing exact reasoning and confidence bounds."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Cost Efficiency: Instant accessibility across teams directly via web browser without specialized hardware."
            p.font.size = Pt(18)
            
    img_path4 = "slide_6_7_extracted/screenshot1_landing.png"
    if os.path.exists(img_path4):
        slide5.shapes.add_picture(img_path4, Inches(0.5), Inches(3.5), width=Inches(8))

    # ---------------- SLIDE 6: Research and References ----------------
    slide6 = prs.slides[5]
    update_team_name(slide6)
    for shape in slide6.shapes:
        if shape.name == 'TextBox 8' and shape.has_text_frame:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "Research & Data Sources:"
            p.font.bold = True
            p.font.size = Pt(20)
            p = tf.add_paragraph()
            p.text = "• Chandrayaan-2 Dual-Frequency Synthetic Aperture Radar (DFSAR) parameters (CPR > 1, DOP < 0.13)."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Chandrayaan-2 CLASS X-Ray Spectrometer and NASA LOLA Altimeter."
            p.font.size = Pt(18)
            p = tf.add_paragraph()
            p.text = "• Craig Reynolds' Flocks, Herds, and Schools: A Distributed Behavioral Model (Boids)."
            p.font.size = Pt(18)

    # ---------------- Delete Slide 7 ----------------
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    if len(slides) > 6:
        xml_slides.remove(slides[6])

    prs.save("SIH_Submission_StormBreaker_Final.pptx")
    print("Saved SIH_Submission_StormBreaker_Final.pptx")

if __name__ == '__main__':
    update_sih_presentation()
