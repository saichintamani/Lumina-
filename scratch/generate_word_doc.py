import os
import requests
import qrcode
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def download_image(url, filename, headers=None):
    if headers is None:
        headers = {'User-Agent': 'Mozilla/5.0'}
    try:
        response = requests.get(url, headers=headers, timeout=10)
        response.raise_for_status()
        with open(filename, 'wb') as f:
            f.write(response.content)
        return filename
    except Exception as e:
        print(f"Error downloading {url}: {e}")
        return None

def generate_qr(data, filename):
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=12,
        border=2,
    )
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img.save(filename)
    return filename

def add_header(doc, text, color=RGBColor(25, 55, 100), font_size=14):
    p = doc.add_heading(level=2)
    run = p.add_run(text)
    run.font.color.rgb = color
    run.font.size = Pt(font_size)
    run.bold = True
    return p

def add_reference(doc, category, text, link):
    p = doc.add_paragraph()
    
    # Custom styled label
    label = p.add_run(f"[{category}] ")
    label.bold = True
    label.font.color.rgb = RGBColor(0, 80, 160)
    
    p.add_run(f"{text} — ")
    
    link_run = p.add_run(link)
    link_run.font.color.rgb = RGBColor(0, 112, 192)
    link_run.underline = True
    p.paragraph_format.space_after = Pt(6)

def main():
    doc = Document()
    
    # MAIN TITLE (Like SIH Template)
    title = doc.add_heading('RESEARCH AND REFERENCES', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    title_run.font.size = Pt(24)
    title_run.bold = True
    
    doc.add_paragraph('Smart India Hackathon 2026 (SIH26166) | Team: LumaInit (StormBreaker)\n').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # -----------------------------------------
    # REFERENCES SECTION
    # -----------------------------------------
    add_header(doc, "PUBLIC DATASETS & RESEARCH")
    
    add_reference(doc, "ISRO DATA", "Chandrayaan-2 DFSAR Radar backscatter & elevation data", "https://pradan.issdc.gov.in")
    add_reference(doc, "ISRO DATA", "Chandrayaan-2 CLASS (X-Ray Spectrometer) elemental abundance maps", "https://pradan.issdc.gov.in")
    add_reference(doc, "NASA DATA", "NASA LRO LOLA Altimeter digital elevation models", "https://pds-geosciences.wustl.edu")
    add_reference(doc, "NASA DATA", "NASA JPL Horizons Ephemeris – real lunar orbital physics", "https://ssd.jpl.nasa.gov/horizons/")
    add_reference(doc, "KAGGLE DATA", "Lumina Synthetic Lunar Surface Matching Dataset", "https://www.kaggle.com/datasets/saichintamaniai/isro-lunar-surface-matching")
    add_reference(doc, "KAGGLE CODE", "Lumina Deep Learning LoFTR Feature Matching Pipeline", "https://www.kaggle.com/code/saichintamaniai/isro-lunar-surface-loftr-feature-matching")
    add_reference(doc, "ALGORITHM", "Craig W. Reynolds, 'Flocks, Herds, and Schools: A Distributed Behavioral Model'", "SIGGRAPH 1987")
    add_reference(doc, "OFFICIAL", "ISRO Official Portal – Bharatiya Antariksh Hackathon problem statement", "https://isro.gov.in")
    
    doc.add_paragraph("\n")
    
    # -----------------------------------------
    # LIVE DEPLOYMENTS & QR CODES
    # -----------------------------------------
    add_header(doc, "OPEN & REPRODUCIBLE PIPELINE")
    subtitle = doc.add_paragraph("Every dataset used is public and every claim in this submission is independently verifiable.")
    subtitle.runs[0].italic = True
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.color.rgb = RGBColor(160, 100, 0)
    
    # Create a 2-column table for GitHub and Vercel
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    
    # Setup cell 1 (GITHUB)
    cell_github = table.cell(0, 0)
    gh_p1 = cell_github.add_paragraph()
    gh_run = gh_p1.add_run("PROJECT REPOSITORY (GITHUB)\n")
    gh_run.bold = True
    gh_run.font.color.rgb = RGBColor(0, 80, 160)
    gh_p1.add_run("https://github.com/saichintamani/Lumina-\n\n")
    gh_p1.add_run("SCAN → GITHUB REPO\n").bold = True
    gh_p1.add_run("Full source code, README & architecture docs")
    gh_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    generate_qr("https://github.com/saichintamani/Lumina-", "gh_qr.png")
    try:
        gh_p1.add_run("\n").add_picture("gh_qr.png", width=Inches(1.5))
    except Exception as e:
        print("Error with GH QR:", e)

    # Setup cell 2 (VERCEL)
    cell_vercel = table.cell(0, 1)
    vc_p1 = cell_vercel.add_paragraph()
    vc_run = vc_p1.add_run("LIVE DEPLOYMENT (VERCEL)\n")
    vc_run.bold = True
    vc_run.font.color.rgb = RGBColor(0, 128, 64)
    vc_p1.add_run("https://lumina-zeta-sand.vercel.app\n\n")
    vc_p1.add_run("SCAN → LIVE DEMO\n").bold = True
    vc_p1.add_run("Try the working 3D prototype instantly, no install")
    vc_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    generate_qr("https://lumina-zeta-sand.vercel.app", "vc_qr.png")
    try:
        vc_p1.add_run("\n").add_picture("vc_qr.png", width=Inches(1.5))
    except Exception as e:
        print("Error with VC QR:", e)

    # Clean up QRs
    for f in ["gh_qr.png", "vc_qr.png"]:
        if os.path.exists(f):
            os.remove(f)

    # Save
    output_path = r"d:\My projects\ISRO\Lumina_Project_References_V2.docx"
    doc.save(output_path)
    print(f"Professional SIH Template Word Doc generated at: {output_path}")

if __name__ == "__main__":
    main()
