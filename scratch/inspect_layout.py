from pptx import Presentation

prs = Presentation("SIH_Submission_StormBreaker.pptx")
print(f"Slide width: {prs.slide_width.inches} inches, Slide height: {prs.slide_height.inches} inches")

for i, slide in enumerate(prs.slides):
    if i == 0: continue # skip title slide for now
    print(f"\nSlide {i+1}:")
    for shape in slide.shapes:
        if shape.name == 'TextBox 8':
            print(f"  TextBox 8: left={shape.left.inches:.2f}, top={shape.top.inches:.2f}, width={shape.width.inches:.2f}, height={shape.height.inches:.2f}")
        elif shape.has_text_frame and shape.text_frame.text.strip():
            print(f"  {shape.name}: text='{shape.text_frame.text[:20]}', left={shape.left.inches:.2f}, top={shape.top.inches:.2f}")
