import sys
from pptx import Presentation

def inspect_pptx(file_path):
    try:
        prs = Presentation(file_path)
        print(f"Presentation: {file_path}")
        print(f"Total Slides: {len(prs.slides)}")
        
        # Print layout names
        print("\nSlide Master Layouts:")
        for idx, l in enumerate(prs.slide_layouts):
            print(f"Layout {idx}: {l.name}")
            
        print("\nSlide Contents:")
        for idx, slide in enumerate(prs.slides):
            print(f"\n--- Slide {idx + 1} ---")
            if slide.shapes.title:
                print(f"Title: {slide.shapes.title.text}")
            else:
                print("Title: [No Title Shape]")
                
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        print(f"  Shape [{shape.name}] Text: {text}")
                elif shape.shape_type == 13: # Picture
                    print(f"  Shape [{shape.name}] is a Picture")
                else:
                    print(f"  Shape [{shape.name}] type: {shape.shape_type}")
    except Exception as e:
        print(f"Error inspecting pptx: {str(e)}")

if __name__ == "__main__":
    file_path = "BAH2026_Submission_Deck.pptx"
    if len(sys.argv) > 1:
        file_path = sys.argv[1]
    inspect_pptx(file_path)
