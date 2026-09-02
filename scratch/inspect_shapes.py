from pptx import Presentation

prs = Presentation("SIH_Submission_StormBreaker.pptx")
for s_idx, slide in enumerate(prs.slides):
    print(f"\n================ SLIDE {s_idx + 1} ================")
    for sh_idx, shape in enumerate(slide.shapes):
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text.replace("\n", " | ")
        print(f"Shape {sh_idx}: Name='{shape.name}', Type={shape.shape_type}, Text='{text}'")
